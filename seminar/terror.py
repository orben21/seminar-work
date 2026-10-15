"""The Global Terrorism Database side: File 3, Part A (which events qualify) and
Step 1 of Part B (aggregating qualifying events to the market-year).

An event counts towards the study only if it passes, in this order:

    window  1998-2020                       Part 0.1; File 3 §3.2
    K1      GTD baseline attributes         true of every GTD record: the GTD
                                            only admits intentional, violent,
                                            sub-national acts (START, 2021)
    K2      all three inclusion criteria    crit1 = crit2 = crit3 = 1
    K3      no definitional doubt           doubtterr = 0; the -9 "not coded"
                                            value is excluded, which is why the
                                            window cannot open before 1998
    K4      enacted, not planned            true of every GTD record; `success`
                                            is recorded but is NOT a filter
    K5      country attribution             the country of occurrence
    K6      single-incident rule            one GTD record = one incident;
                                            coordinated attacks linked through
                                            `related` are counted individually
    floor   nkill >= 1 OR nwound >= 1 OR property = 1        File 3 §4

These are the study's definitions (File 3 §3-§4), so they are coded here with
their source instead of being switches in config.py.

The GTD ships as a 104 MB xlsx with ~210,000 events. It is converted once into
a parquet cache of the fields File 3 names; every later run loads that in under
a second.
"""

from __future__ import annotations

import math
import re
import time
import unicodedata
from pathlib import Path
from typing import Dict, Iterable, List, Optional, Tuple

import numpy as np
import pandas as pd

from . import countries, paths


class GTDError(Exception):
    """Raised when the GTD file cannot be read or looks wrong."""


# Only these of the 135 GTD columns are kept: the fields named in plan §5.2
# (qualification), §5.3-§5.4 (severity traits) and §5.8 (supplementary flags).
# Attack type, weapon, perpetrator and hostage fields are deliberately absent -
# the study measures damage to a market, not the tactics of the event (§5.4).
KEEP_COLUMNS = [
    "eventid", "iyear", "extended", "resolution",                  # K6, dates
    "country", "country_txt", "city", "latitude", "longitude",      # K5, flags
    "crit1", "crit2", "crit3", "doubtterr", "alternative", "related",  # K2, K3
    "success", "nkill", "nwound", "property", "propextent",         # traits
    "suicide", "targtype1", "targsubtype1", "natlty1",              # flags
]
NUMERIC_COLUMNS = [
    "eventid", "iyear", "extended", "country", "latitude", "longitude",
    "crit1", "crit2", "crit3", "doubtterr", "alternative", "success",
    "nkill", "nwound", "property", "propextent", "suicide", "targtype1",
    "targsubtype1", "natlty1",
]
TEXT_COLUMNS = ["country_txt", "city", "related", "resolution"]

# The qualification funnel of plan Figure 11, in order.
FUNNEL_STAGES = [
    ("window", "All GTD events in window"),
    ("K1", "K1 GTD baseline attributes"),
    ("K2", "K2 all three criteria (crit1-3 = 1)"),
    ("K3", "K3 no definitional doubt (doubtterr = 0)"),
    ("K4-K6", "K4-K6 enacted, attributed, single incident"),
    ("floor", "Magnitude floor (any recorded loss)"),
]

# File 3 §7 Step 1: property damage score per event, from the propextent band.
# GTD codes 1 = catastrophic, 2 = major, 3 = minor, 4 = unknown. The weights
# are ordinal, not monetary (File 3 §8). Unknown is recoded to minor and
# flagged (File 3 §5.1) - the conservative direction.
PROPEXTENT_WEIGHT = {1: 3.0, 2: 2.0, 3: 1.0, 4: 1.0}
PROPEXTENT_LABELS = {1: "catastrophic", 2: "major", 3: "minor",
                     4: "unknown (recoded to minor)"}

# GTD target codes used by the §5.8 flags (START, 2021, codebook §IV).
TARGET_BUSINESS = 1                 # targtype1 = Business
TARGET_UTILITIES = 21               # targtype1 = Utilities (power, pipelines)
SUBTYPE_GAS_OIL_ELECTRIC = 1        # targsubtype1 under Business

FLAG_COLUMNS = [
    "flag_suicide", "flag_capital", "flag_capital_100km", "flag_top10_city",
    "flag_property", "flag_foreign_target", "flag_economic_target",
    "flag_energy_target",
]
# Flags that need data/reference/cities.csv; unknown (NaN) without it.
CITY_FLAGS = ["flag_capital", "flag_capital_100km", "flag_top10_city"]

COMPONENTS = ["NA", "NK", "NI", "PD"]
CAPITAL_RADIUS_KM = 100.0           # Ahmad et al. (2022): within 100 km


# --------------------------------------------------------------------------
# one-time conversion
# --------------------------------------------------------------------------

def convert_to_parquet(xlsx_path: Path,
                       out_path: Path = paths.GTD_EVENTS_PARQUET,
                       sheet_name: str = "Data",
                       chunk_size: int = 25_000) -> pd.DataFrame:
    """Stream the GTD xlsx into a parquet cache of just the fields we need.

    Streaming (openpyxl read_only) keeps memory flat; the whole sheet never
    sits in RAM at once. Events whose country cannot be resolved are KEPT with
    an empty iso3, so the Q1 audit can report them rather than lose them.
    """
    from openpyxl import load_workbook

    xlsx_path = Path(xlsx_path)
    if not xlsx_path.exists():
        raise GTDError(f"GTD file not found: {xlsx_path}\n"
                       f"Check GTD_XLSX in config.py.")

    print(f"Reading {xlsx_path.name} ({xlsx_path.stat().st_size / 1e6:.0f} MB)")
    print("This takes a few minutes, but only has to happen once.")
    started = time.time()

    workbook = load_workbook(xlsx_path, read_only=True, data_only=True)
    if sheet_name not in workbook.sheetnames:
        raise GTDError(f"sheet '{sheet_name}' not found. "
                       f"Sheets: {workbook.sheetnames}")
    rows = workbook[sheet_name].iter_rows(values_only=True)

    header = [str(h).strip() if h is not None else "" for h in next(rows)]
    position_of = {name: i for i, name in enumerate(header)}
    missing = [c for c in KEEP_COLUMNS if c not in position_of]
    if missing:
        raise GTDError(f"these expected GTD columns are absent: {missing}\n"
                       f"Is this really the GTD distribution file?")
    positions = [position_of[c] for c in KEEP_COLUMNS]

    chunks: List[pd.DataFrame] = []
    buffer: List[tuple] = []
    n_rows = 0
    for row in rows:
        width = len(row)
        buffer.append(tuple(row[p] if p < width else None for p in positions))
        n_rows += 1
        if len(buffer) >= chunk_size:
            chunks.append(pd.DataFrame(buffer, columns=KEEP_COLUMNS))
            buffer = []
            print(f"  {n_rows:>7,} events   ({time.time() - started:.0f}s)",
                  flush=True)
    if buffer:
        chunks.append(pd.DataFrame(buffer, columns=KEEP_COLUMNS))
    workbook.close()

    events = pd.concat(chunks, ignore_index=True)
    print(f"  {len(events):,} events read in {time.time() - started:.0f}s")

    for column in NUMERIC_COLUMNS:
        events[column] = pd.to_numeric(events[column], errors="coerce")
    for column in TEXT_COLUMNS:
        text = events[column].astype("string").str.strip()
        events[column] = text.where(text != "", pd.NA)

    events = events.dropna(subset=["iyear", "eventid"])
    events["iyear"] = events["iyear"].astype(int)
    events["eventid"] = events["eventid"].astype("int64")
    events["iso3"] = countries.resolve(events["country_txt"], source="GTD").values

    unresolved = events["iso3"].isna()
    if unresolved.any():
        names = sorted(events.loc[unresolved, "country_txt"].dropna().unique())
        print(f"  note: {int(unresolved.sum()):,} events are in locations that "
              f"do not resolve to a country ({names[:8]}); they stay in the "
              f"cache and fail K5 / Q1 in the audit")

    out_path.parent.mkdir(parents=True, exist_ok=True)
    events.to_parquet(out_path, index=False)
    print(f"  cached -> {out_path}")
    return events


def load_events(cache_path: Path = paths.GTD_EVENTS_PARQUET) -> pd.DataFrame:
    """Load the cached events. Tells you to build the cache if it is missing."""
    cache_path = Path(cache_path)
    if not cache_path.exists():
        raise GTDError(f"no GTD cache at {cache_path}\n"
                       f"Run 1_build_terror_cache.py once to create it.")
    events = pd.read_parquet(cache_path)
    absent = [c for c in KEEP_COLUMNS + ["iso3"] if c not in events.columns]
    if absent:
        raise GTDError(f"the GTD cache predates this version of the code and "
                       f"lacks {absent}.\nRun 1_build_terror_cache.py again.")
    return events


# --------------------------------------------------------------------------
# File 3 Part A: which events qualify
# --------------------------------------------------------------------------

def stage_masks(events: pd.DataFrame,
                window: Tuple[int, int]) -> Dict[str, pd.Series]:
    """Cumulative pass/fail of every event at every stage of FUNNEL_STAGES."""
    first, last = window
    in_window = events["iyear"].between(first, last)
    k1 = in_window                              # true of every GTD record
    k2 = k1 & (events["crit1"] == 1) & (events["crit2"] == 1) \
        & (events["crit3"] == 1)
    k3 = k2 & (events["doubtterr"] == 0)        # -9 (not coded) fails too
    k4_k6 = (k3 & events["iso3"].notna()        # K5: a country of occurrence
             & ~events["eventid"].duplicated())  # K6: one record, one incident
    floor = k4_k6 & ((events["nkill"] >= 1) | (events["nwound"] >= 1)
                     | (events["property"] == 1))
    return {"window": in_window, "K1": k1, "K2": k2, "K3": k3,
            "K4-K6": k4_k6, "floor": floor}


def qualifying_events(events: pd.DataFrame,
                      window: Tuple[int, int]) -> pd.DataFrame:
    """Events that pass the window, K1-K6 and the magnitude floor."""
    return events[stage_masks(events, window)["floor"]].copy()


def event_funnel(events: pd.DataFrame, window: Tuple[int, int],
                 population: str) -> pd.DataFrame:
    """Figure 11: how many events survive each stage, and what each removes."""
    masks = stage_masks(events, window)
    start = int(masks["window"].sum())
    rows, previous = [], start
    for key, label in FUNNEL_STAGES:
        remaining = int(masks[key].sum())
        rows.append({
            "population": population, "stage": key, "label": label,
            "events": remaining,
            "pct_of_window": remaining / start * 100 if start else np.nan,
            "removed": previous - remaining,
            "pct_removed_at_stage": ((previous - remaining) / previous * 100
                                     if previous else np.nan),
        })
        previous = remaining
    return pd.DataFrame(rows)


# --------------------------------------------------------------------------
# severity traits (§5.3-§5.4) and supplementary flags (§5.8)
# --------------------------------------------------------------------------

def add_traits(qualifying: pd.DataFrame) -> pd.DataFrame:
    """Per-event severity traits and flags. Nothing here changes which events
    count - that was settled by stage_masks."""
    q = qualifying.copy()
    damaged = q["property"] == 1
    band = q["propextent"]

    # PD score (File 3 §7 Step 1). property = 1 with a blank band is treated
    # like band 4: damage occurred, extent unknown -> minor, flagged.
    q["pd_score"] = np.where(damaged, band.map(PROPEXTENT_WEIGHT).fillna(1.0),
                             0.0)
    q["pd_unknown_recoded"] = damaged & (band.isna() | (band == 4))
    # property = -9: unknown whether any damage occurred. Scored 0 (no damage
    # is asserted) and counted, so the reader can see how often it happens.
    q["property_unknown"] = q["property"] == -9

    # Blank is not zero (File 3 §10): these stay NaN in nkill / nwound.
    q["nkill_blank"] = q["nkill"].isna()
    q["nwound_blank"] = q["nwound"].isna()

    q["flag_suicide"] = q["suicide"] == 1
    q["flag_property"] = damaged
    # natlty1 uses the same numeric codes as `country`, so "differs from the
    # country of occurrence" is an exact comparison (File 3 §9.1).
    q["flag_foreign_target"] = q["natlty1"].notna() & (q["natlty1"]
                                                       != q["country"])
    q["flag_economic_target"] = q["targtype1"].isin([TARGET_BUSINESS,
                                                     TARGET_UTILITIES])
    q["flag_energy_target"] = (
        ((q["targtype1"] == TARGET_BUSINESS)
         & (q["targsubtype1"] == SUBTYPE_GAS_OIL_ELECTRIC))
        | (q["targtype1"] == TARGET_UTILITIES))
    for flag in CITY_FLAGS:
        if flag not in q.columns:
            q[flag] = np.nan
    return q


def _normalise_place(name) -> str:
    """Lowercase ASCII with punctuation removed, for matching city names."""
    if name is None or (isinstance(name, float) and math.isnan(name)):
        return ""
    text = unicodedata.normalize("NFKD", str(name))
    text = text.encode("ascii", "ignore").decode("ascii").lower()
    text = re.sub(r"[^a-z0-9 ]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def _haversine_km(lat1, lon1, lat2, lon2):
    lat1, lon1, lat2, lon2 = map(np.radians, (lat1, lon1, lat2, lon2))
    a = (np.sin((lat2 - lat1) / 2) ** 2
         + np.cos(lat1) * np.cos(lat2) * np.sin((lon2 - lon1) / 2) ** 2)
    return 6371.0 * 2 * np.arcsin(np.sqrt(a))


def add_city_flags(q: pd.DataFrame, cities: Optional[pd.DataFrame],
                   capital_overrides: Optional[Dict[str, str]] = None
                   ) -> Tuple[pd.DataFrame, List[str]]:
    """The capital, 100 km and top-ten-city flags of File 3 §9.1.

    `cities` is data/reference/cities.csv (built by 0_fetch_data.py from
    GeoNames): per country its capital and its ten most populous other cities.
    Names are matched on the GTD `city` field, distance on latitude/longitude.
    A country with no known capital keeps these flags empty (unknown), and is
    returned in the list of notes.
    """
    notes: List[str] = []
    q = q.copy()
    for flag in CITY_FLAGS:
        q[flag] = np.nan
    if cities is None or cities.empty:
        notes.append("data/reference/cities.csv is missing, so the capital and "
                     "top-ten-city flags are empty. Run 0_fetch_data.py.")
        return q, notes

    overrides = {k.upper(): v for k, v in (capital_overrides or {}).items()}
    city_norm = q["city"].map(_normalise_place)

    for iso3, rows in q.groupby("iso3").groups.items():
        own = cities[cities["iso3"] == iso3]
        if own.empty:
            notes.append(f"{iso3}: no cities in the reference file")
            continue
        capital = own[own["is_capital"] == 1]
        if iso3 in overrides:
            wanted = _normalise_place(overrides[iso3])
            capital = own[own["names"].map(
                lambda names: wanted in names.split("|"))].head(1)
            if capital.empty:
                notes.append(f"{iso3}: CAPITAL_OVERRIDES names "
                             f"'{overrides[iso3]}', which is not among its "
                             f"cities in the reference file")
        top10 = own[own["is_capital"] != 1].nlargest(10, "population")
        top10_names = {n for names in top10["names"] for n in names.split("|")
                       if n}
        index = list(rows)
        q.loc[index, "flag_top10_city"] = city_norm.loc[index].isin(
            top10_names).astype(float)
        if capital.empty:
            notes.append(f"{iso3}: no capital known, so its capital flags are "
                         f"empty (set CAPITAL_OVERRIDES in config.py)")
            continue
        cap = capital.iloc[0]
        capital_names = {n for n in str(cap["names"]).split("|") if n}
        q.loc[index, "flag_capital"] = city_norm.loc[index].isin(
            capital_names).astype(float)
        distance = _haversine_km(q.loc[index, "latitude"].astype(float),
                                 q.loc[index, "longitude"].astype(float),
                                 float(cap["latitude"]),
                                 float(cap["longitude"]))
        q.loc[index, "flag_capital_100km"] = np.where(
            distance.isna(), np.nan, (distance <= CAPITAL_RADIUS_KM) * 1.0)
    return q, notes


def load_cities(path: Path = paths.CITIES_CSV) -> Optional[pd.DataFrame]:
    """The GeoNames-derived cities reference, or None if it is not there."""
    path = Path(path)
    if not path.exists():
        return None
    cities = pd.read_csv(path, encoding="utf-8-sig", keep_default_na=False,
                         na_values=[""])
    cities["names"] = cities["names"].fillna("").astype(str)
    return cities


# --------------------------------------------------------------------------
# File 3 Part B, Step 1: aggregate to the market-year
# --------------------------------------------------------------------------

def aggregate_market_years(q: pd.DataFrame, iso3_codes: Iterable[str],
                           window: Tuple[int, int]) -> pd.DataFrame:
    """NA, NK, NI, PD and the flag counts for every market-year of the window.

    A market-year with no qualifying event is a genuine zero and stays in the
    panel (File 3 §7 Step 2). NK / NI sum the events whose count is known; a
    market-year whose every qualifying event has a blank count is MISSING on
    that component, never zero (File 3 §10), and the number of blank events
    is reported alongside so the reader can see how much rests on them.
    Extended incidents are counted in `iyear`, the year they began (§7 Step 1).
    """
    iso3_codes = sorted(set(iso3_codes))
    q = q[q["iso3"].isin(iso3_codes)]
    keys = ["iso3", "iyear"]
    grouped = q.groupby(keys)

    table = grouped.agg(
        NA=("eventid", "count"),
        PD=("pd_score", "sum"),
        nk_blank_events=("nkill_blank", "sum"),
        ni_blank_events=("nwound_blank", "sum"),
        pd_unknown_recoded=("pd_unknown_recoded", "sum"),
        property_unknown_events=("property_unknown", "sum"),
        unsuccessful_events=("success", lambda s: int((s == 0).sum())),
        extended_events=("extended", lambda s: int((s == 1).sum())),
    )
    table["NK"] = grouped["nkill"].sum(min_count=1)
    table["NI"] = grouped["nwound"].sum(min_count=1)
    for flag in FLAG_COLUMNS:
        column = f"n_{flag.replace('flag_', '')}"
        table[column] = grouped[flag].sum(min_count=1)
    table = table.reset_index().rename(columns={"iyear": "year"})

    first, last = window
    grid = pd.DataFrame(
        [(code, year) for code in iso3_codes
         for year in range(first, last + 1)], columns=["iso3", "year"])
    merged = grid.merge(table, on=["iso3", "year"], how="left")

    no_events = merged["NA"].isna()
    count_columns = [c for c in merged.columns if c not in ("iso3", "year")]
    merged.loc[no_events, count_columns] = 0.0
    ordered = (["iso3", "year"] + COMPONENTS
               + [c for c in count_columns if c not in COMPONENTS])
    merged = merged[ordered]
    for column in ordered[2:]:
        merged[column] = merged[column].astype(float)
    return merged.sort_values(["iso3", "year"]).reset_index(drop=True)


def export_event_table(q: pd.DataFrame) -> pd.DataFrame:
    """The event-level Stage 1 export: named fields, derived traits, flags."""
    derived = ["iso3", "pd_score", "pd_unknown_recoded", "property_unknown",
               "nkill_blank", "nwound_blank"] + FLAG_COLUMNS
    columns = [c for c in KEEP_COLUMNS + derived if c in q.columns]
    return q[columns].sort_values(["iso3", "iyear", "eventid"])


def describe(events: pd.DataFrame, window: Tuple[int, int]) -> str:
    """A short readable summary of the cache, for the build script."""
    first, last = window
    inside = events[events["iyear"].between(first, last)]
    doubt = inside["doubtterr"].value_counts().to_dict()
    lines = [
        f"events in cache      : {len(events):,} "
        f"({int(events['iyear'].min())}-{int(events['iyear'].max())})",
        f"events {first}-{last}  : {len(inside):,}",
        f"countries in window  : {inside['country_txt'].nunique()} GTD names",
        f"doubtterr in window  : " + ", ".join(
            f"{int(k)}: {v:,}" for k, v in sorted(doubt.items())),
    ]
    if last > int(events["iyear"].max()):
        lines.append(f"WARNING: the window ends in {last} but this GTD release "
                     f"stops in {int(events['iyear'].max())}")
    return "\n".join("  " + line for line in lines)
